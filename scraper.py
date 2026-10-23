#!/usr/bin/env python3
"""
Healthcare Website Booking Feature Scraper

Scrapes healthcare company websites to determine if they offer
online appointment/booking functionality.
"""

import re
import time
import signal
import logging
from collections import deque
from typing import Optional, Tuple, List, Set
from urllib.parse import urlparse, urljoin, urldefrag
from dataclasses import dataclass
from enum import Enum

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
import urllib3

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

logging.basicConfig(
    level=logging.INFO,
    format='[%(asctime)s] %(levelname)s - %(message)s',
    datefmt='%H:%M:%S'
)
logger = logging.getLogger(__name__)


# Local workbook (Google Sheets API removed for speed).
LOCAL_SAVE_EVERY = 10  # flush to disk every N cell writes

WEBSITE_HEADER = "Website"
BOOKING_HEADER = "Booking Feature"

REQUEST_TIMEOUT = 15          # read timeout (connect is shorter; see _get_once)
CONNECT_TIMEOUT = 6
REQUEST_DELAY = 2.5           # pause between sites (reduces IP / bot blocks)
PAGE_FETCH_DELAY = 0.35       # pause between pages on the same site

MAX_NETWORK_RETRIES = 1
MAX_HTTP_RETRIES = 1
# Full fetch rounds for flaky hosts (user often needs several browser refreshes).
FLAKY_FETCH_ROUNDS = 3
FLAKY_FETCH_BACKOFF = 0.9
FLAKY_READ_TIMEOUT = 10       # shorter per-attempt timeout during refresh rounds
MAX_FETCH_SECONDS = 45        # hard wall-clock budget per site fetch
MAX_PAGES_PER_SITE = 4
JS_FOLLOW_PAGES = 0
JS_RENDER_TIMEOUT_MS = 15000
JS_SETTLE_SECONDS = 2
UC_RENDER_TIMEOUT_SEC = 35
ALTERNATE_DISCOVERY = True
# Playwright + SeleniumBase UC (each can spawn Chromium ≈200–500MB+).
# Turned OFF automatically with --light (needed for multi-process runs).
ENABLE_JS_RENDER = True
# Cap stored HTML so BeautifulSoup does not hold multi‑MB pages in RAM.
MAX_HTML_BYTES = 1_500_000


def apply_light_mode() -> None:
    """Low-RAM settings for running several scrapers in parallel."""
    global LOCAL_SAVE_EVERY, REQUEST_DELAY, PAGE_FETCH_DELAY
    global FLAKY_FETCH_ROUNDS, MAX_FETCH_SECONDS, MAX_PAGES_PER_SITE
    global ALTERNATE_DISCOVERY, ENABLE_JS_RENDER, MAX_HTML_BYTES

    LOCAL_SAVE_EVERY = 40
    REQUEST_DELAY = 3.0
    PAGE_FETCH_DELAY = 0.2
    FLAKY_FETCH_ROUNDS = 2
    MAX_FETCH_SECONDS = 25
    MAX_PAGES_PER_SITE = 2
    ALTERNATE_DISCOVERY = False
    ENABLE_JS_RENDER = False
    MAX_HTML_BYTES = 800_000
    logger.info(
        "LIGHT MODE: no Playwright/UC browsers, 2 pages/site, "
        "no alternate discovery, save every 40 writes"
    )

# Strong interstitial markers — ONLY strings that appear on bot walls, not on
# normal Cloudflare-backed sites (those embed challenge-platform scripts too).
CHALLENGE_MARKERS_STRONG = [
    'sgcaptcha',
    '/.well-known/captcha',
    '/.well-known/sgcaptcha',
    'cf-browser-verification',
    'challenge-running',
    'captcha-delivery.com',
    'you have been blocked',
    'browsing and clicking at a speed much faster',
    'please verify you are a human',
    'verify you are human',
    'are you a robot',
    'enable javascript and cookies to continue',
    'checking your browser before accessing',
    'checking if the site connection is secure',
    'just a moment...',
    'attention required! | cloudflare',
    'attention required | cloudflare',
    'ddos protection by cloudflare',
]

# Only on SMALL pages — CF storefronts also load these scripts.
CHALLENGE_MARKERS_WEAK = [
    'hcaptcha.com/1/api.js',
    'challenges.cloudflare.com/turnstile',
    'challenges.cloudflare.com',
    'challenge-platform',
    'cf-challenge',
    '__cf_chl',
    'turnstile/v0/api.js',
    'perimeterx',
    'datadome',
    'distil networks',
    'bot detection',
    'access denied',
]

# Minimum HTML size to treat as a real page (not a CF interstitial).
MIN_REAL_PAGE_BYTES = 12000

# Skip finished cells on resume. ERROR retry is OFF for now so a full pass
# finishes blank rows first; uncomment the first DONE_VALUES (and comment the
# second) when ready to re-scan ERROR rows.
# DONE_VALUES = {'YES', 'NO', 'SKIP'}  # ERROR omitted → retry ERROR on resume
DONE_VALUES = {'YES', 'NO', 'SKIP', 'ERROR'}


def _strip_www(host: str) -> str:
    """Remove a leading www. prefix only (never use str.lstrip('www.'))."""
    host = (host or '').lower()
    return host[4:] if host.startswith('www.') else host


BOOKING_KEYWORDS = [
    # English — medical / general
    'book appointment', 'book an appointment', 'schedule appointment',
    'schedule an appointment', 'make an appointment',
    'request an appointment', 'appointment request', 'book online',
    'schedule online', 'book now', 'schedule now', 'book a visit',
    'schedule a visit', 'schedule a consultation', 'book a consultation',
    'online booking', 'online scheduling', 'online appointment',
    'patient scheduling', 'appointment booking', 'find a time',
    'choose a time', 'select a date', 'reserve a slot', 'request appointment',
    'schedule your appointment', 'book your appointment', 'make appointment',
    'request visit', 'schedule visit', 'book visit', 'appointment form',
    'book a meeting', 'schedule a meeting', 'schedule a call',
    'book a call', 'schedule consultation', 'request a consultation',
    'make a booking', 'book a discovery', 'book discovery', 'online bookings',
    'book here', 'booking online', 'reserve online', 'reserve now',
    'book your session', 'schedule a session', 'book a session',
    'book a consult', 'book a consultation', 'book an assessment',
    'schedule an assessment', 'schedule a screening', 'book a screening',
    # Telehealth / virtual care (common on clinic & therapy sites)
    'virtual visit', 'schedule a virtual visit', 'book a virtual visit',
    'telehealth appointment', 'schedule telehealth', 'book telehealth',
    'schedule a telehealth', 'virtual appointment', 'online visit',
    'start therapy', 'begin therapy', 'start treatment', 'get care now',
    'new patient appointment', 'schedule new patient',
    'same-day appointment', 'same day appointment', 'walk-in appointment',
    'patient intake', 'complete intake', 'intake form', 'online intake',
    'reserve your spot', 'reserve a spot',
    # Senior living / assisted living tours
    'schedule a tour', 'schedule tour', 'book a tour', 'book tour',
    'request a tour', 'request tour', 'schedule your tour', 'arrange a tour',
    # Hotels / resorts / stays
    'book a room', 'book your room', 'book a stay', 'book your stay',
    'reserve a room', 'reserve your room', 'reserve a stay',
    'check availability', 'check rates', 'book rates', 'book hotel',
    'book this hotel', 'make a reservation', 'online reservation',
    'room reservation', 'hotel reservation', 'reserve online',
    'book accommodations', 'book accommodation', 'book your getaway',
    'check in online', 'modify reservation', 'book direct',
    # Spa / salon / wellness / beauty
    'book a treatment', 'book treatment', 'book a service', 'book service',
    'book a spa', 'book spa', 'book a facial', 'book facial',
    'book a massage', 'book massage', 'book a wax', 'book waxing',
    'book a haircut', 'book haircut', 'book a color', 'book color',
    'book a manicure', 'book manicure', 'book a pedicure', 'book pedicure',
    'book lashes', 'book brows', 'book injectables', 'book botox',
    'book a package', 'book package', 'book spa day', 'book a spa day',
    'schedule a treatment', 'schedule treatment', 'book & shop',
    'book and shop', 'book/shop', 'spa booking', 'salon booking',
    'book your treatment', 'book your service', 'reserve a treatment',
    # Dining / experiences (resort spas)
    'book a table', 'reserve a table', 'make a reservation',
    'book experience', 'book an experience', 'book a tee time',
    # German
    'termin buchen', 'termin vereinbaren', 'online termin', 'termin online',
    'jetzt buchen', 'online buchen', 'termin anfragen', 'termin reservieren',
    'terminbuchung', 'online-termin', 'termin buchen online',
    'arzttermin', 'sprechstunde buchen', 'termin online vereinbaren',
    'zimmer buchen', 'jetzt reservieren', 'online reservieren',
    # French
    'prendre rendez-vous', 'prendre rdv', 'rendez-vous en ligne',
    'réserver en ligne', 'reserver en ligne', 'prendre un rendez-vous',
    'rdv en ligne', 'réservation en ligne', 'reservation en ligne',
    'prendre mon rendez-vous', 'réserver une chambre', 'réserver maintenant',
    # Spanish
    'pedir cita', 'reservar cita', 'cita online', 'cita en línea',
    'sacar cita', 'agendar cita', 'reservar turno', 'reservar hora',
    'pide cita', 'solicitar cita', 'cita médica online',
    'reservar habitación', 'reservar ahora', 'hacer una reserva',
    # Italian
    'prenota online', 'prenota un appuntamento', 'prenotazione online',
    'prenota ora', 'fissa un appuntamento', 'prenota visita',
    'prenota camera', 'prenota soggiorno',
    # Dutch
    'afspraak maken', 'online afspraak', 'afspraak boeken',
    'maak een afspraak', 'plan een afspraak', 'kamer boeken',
    # Portuguese
    'agendar consulta', 'marcar consulta', 'agendamento online',
    'marque sua consulta', 'reservar consulta', 'reservar quarto',
    # Polish
    'umów wizytę', 'umow wizyte', 'rezerwacja online', 'zarezerwuj wizytę',
    'zarezerwuj pokój',
]

# Strong online CTA phrases — enough alone for YES when found in links/buttons.
STRONG_ONLINE_CTAS = [
    'book appointment', 'book an appointment', 'schedule appointment',
    'schedule an appointment', 'book online',
    'schedule online', 'request an appointment', 'request appointment',
    'make an appointment', 'schedule your free appointment',
    'book now', 'schedule now', 'online booking', 'online scheduling',
    'book a visit', 'schedule a visit', 'appointment request',
    'schedule a tour', 'schedule tour', 'book a tour', 'request a tour',
    'book a consult', 'book a consultation', 'virtual visit',
    'schedule a virtual visit', 'telehealth appointment', 'get care now',
    'start therapy', 'begin therapy', 'new patient appointment',
    'same-day appointment', 'same day appointment', 'online intake',
    'make a booking', 'book a call', 'book a discovery', 'book here',
    'reserve now', 'make a reservation', 'online reservation',
    'book a room', 'book your stay', 'book a stay', 'reserve a room',
    'check availability', 'check rates', 'book direct',
    'book a treatment', 'book treatment', 'book a service',
    'book a spa', 'book spa day', 'book & shop', 'book and shop',
    'book a table', 'reserve a table', 'spa booking', 'salon booking',
    # DE / FR / ES / IT / NL
    'termin buchen', 'termin vereinbaren', 'online termin', 'jetzt buchen',
    'online buchen', 'terminbuchung', 'online-termin', 'zimmer buchen',
    'prendre rendez-vous', 'prendre rdv', 'rendez-vous en ligne',
    'réserver en ligne', 'rdv en ligne', 'réserver maintenant',
    'pedir cita', 'reservar cita', 'cita online', 'agendar cita',
    'reservar ahora', 'hacer una reserva',
    'prenota online', 'prenota ora', 'prenotazione online',
    'afspraak maken', 'online afspraak', 'afspraak boeken',
    'agendar consulta', 'marcar consulta', 'agendamento online',
]

# Short link labels that count as YES (exact-ish match on trimmed link text).
STRONG_LINK_LABELS = {
    'book', 'bookings', 'booking', 'book now', 'book online',
    'book appointment', 'book an appointment', 'schedule', 'schedules',
    'appointments', 'appointment', 'request appointment',
    'request an appointment', 'schedule appointment',
    'schedule your free appointment', 'schedule free appointment',
    'make a booking', 'make a booking!', 'reserve', 'reserve now',
    'reservations', 'reservation', 'book a room', 'book stay',
    'book your stay', 'check availability', 'check rates', 'book rates',
    'book treatment', 'book a treatment', 'book service', 'book & shop',
    'book and shop', 'book/shop',
    'schedule tour', 'schedule a tour', 'book a tour', 'book tour',
    'request a tour', 'request tour', 'take a tour', 'arrange a tour',
    'virtual visit', 'telehealth', 'get care now', 'start therapy',
    'begin therapy', 'online intake', 'intake form',
    # DE / FR / ES / IT / NL
    'termin', 'termine', 'termin buchen', 'jetzt buchen', 'online buchen',
    'buchen', 'reservieren',
    'rendez-vous', 'rdv', 'réserver', 'reserver', 'prendre rdv',
    'cita', 'reservar', 'agendar',
    'prenota', 'prenotazione',
    'afspraak', 'afspraak maken',
    'marcar',
}

# URL path patterns that strongly suggest a booking page (not just a portal).
BOOKING_URL_PATTERNS = [
    r'/appointments?(?:/|$|\?)',
    r'/book(?:ing|ings|-online|-now|-a-[a-z0-9-]+)?(?:/|$|\?)',
    r'/schedule(?:-online|-appointment|s)?(?:/|$|\?)',
    r'/scheduling(?:/|$|\?)',
    r'/request-appointment(?:/|$|\?)',
    r'/request-an-appointment(?:/|$|\?)',
    r'/appointment-request(?:/|$|\?)',
    r'/make-(?:an-)?appointment(?:/|$|\?)',
    r'/make-a-booking(?:/|$|\?)',
    r'/new-appointment(?:/|$|\?)',
    r'/visit/schedule(?:/|$|\?)',
    r'/book-an-appointment(?:/|$|\?)',
    r'/online-booking(?:/|$|\?)',
    r'/schedule/request(?:/|$|\?)',
    # Hotels / spa / salon
    r'/reservations?(?:/|$|\?)',
    r'/reserve(?:/|$|\?)',
    r'/booking(?:-engine|s)?(?:/|$|\?)',
    r'/check-availability(?:/|$|\?)',
    r'/check-rates(?:/|$|\?)',
    r'/book-now(?:/|$|\?)',
    r'/book-a-(?:room|stay|treatment|service|table)(?:/|$|\?)',
    r'/spa-booking(?:/|$|\?)',
    r'/salon-booking(?:/|$|\?)',
    r'/book-and-shop(?:/|$|\?)',
    r'/book-shop(?:/|$|\?)',
    # Multilingual path segments
    r'/termin(?:e|buchung|-buchen|-vereinbaren)?(?:/|$|\?)',
    r'/online-termin(?:/|$|\?)',
    r'/rendez[-_]?vous(?:/|$|\?)',
    r'/prendre[-_]?rendez[-_]?vous(?:/|$|\?)',
    r'/cita(?:s)?(?:/|$|\?)',
    r'/reservar(?:-cita)?(?:/|$|\?)',
    r'/prenota(?:zione)?(?:/|$|\?)',
    r'/afspraak(?:/|$|\?)',
    r'/agendar(?:/|$|\?)',
    r'/consulta(?:/|$|\?)',
]

# Host prefixes / names that indicate a booking subdomain.
BOOKING_HOST_PATTERNS = [
    r'^bookings?\.',
    r'^schedule\.',
    r'^scheduling\.',
    r'^appointments?\.',
    r'^reserve\.',
    r'^reservations?\.',
    r'^book\.',
    r'^reservieren\.',
    r'^termin\.',
    r'^termine\.',
    r'^rdv\.',
    r'^cita\.',
    r'\.bookings?\.',
]

# Pages worth following for a deeper look (not automatic YES).
FOLLOW_LINK_PATTERNS = [
    r'appointment', r'booking', r'\bbook\b', r'schedule', r'scheduling',
    r'patient.?portal', r'\bportal\b', r'contact', r'providers?',
    r'doctors?', r'services?', r'request.?visit', r'new.?patient',
    r'make.?a.?booking', r'book.?online',
    r'reservation', r'reserve', r'check.?availability', r'check.?rates',
    r'book.?a.?room', r'book.?your.?stay', r'book.?treatment',
    r'book.?spa', r'book.?and.?shop', r'spa.?booking', r'salon.?booking',
    # Multilingual follow targets
    r'termin', r'buchen', r'reservieren', r'kontakt',
    r'rendez.?vous', r'\brdv\b', r'contactez',
    r'\bcita\b', r'reservar', r'agendar', r'contacto',
    r'prenota', r'appuntamento',
    r'afspraak', r'agendar', r'consulta',
]

# Known booking platforms + AU healthcare schedulers + spa/salon + EHR portals.
# NOTE: Do NOT include Wix bookings-service / BookOnlineViewerWidget /
# bookings.wixapps.net — Wix embeds those in platform JS on sites that
# do not actually offer online booking (false YES).
THIRD_PARTY_BOOKING_DOMAINS = [
    # US healthcare directories / patient scheduling / EHR / therapy
    'zocdoc.com', 'healthgrades.com', 'docasap.com', 'phreesia.com',
    'solutionreach.com', 'athenahealth.com', 'mychart.org', 'mychart.com',
    'doctor.com', 'vitals.com', 'webpt.com', 'nextech.com', 'modmed.com',
    'drchrono.com', 'practicefusion.com', 'kareo.com', 'advancedmd.com',
    'epic.com', 'followmyhealth.com', 'healow.com', 'patientpop.com',
    'demandforce.com', 'lumahealth.io', 'relatient.net', 'doximity.com',
    'amwell.com', 'teladoc.com', 'mdlive.com', 'sessionhealth.com',
    'simplepractice.com', 'theranest.com', 'therapyappointment.com',
    'psychologytoday.com', 'alma.com', 'headway.co', 'growtherapy.com',
    'intakeq.com', 'sprucehealth.com', 'hint.com', 'elationemr.com',
    'charmhealth.com', 'officeally.com', 'rxnt.com', 'ecwcloud.com',
    'eclinicalworks.com', 'nextgen.com', 'allscripts.com', 'cerner.com',
    'bridgepatientportal.com', 'intelichart.com', 'medfusion.net',
    'patientgateway.org', 'mychart.healthcare',
    'patientportal.athenahealth.com', 'patient.athenahealth.com',
    'portal.advancedmd.com', 'secure.followmyhealth.com',
    'patientaccess.com', 'nexhealth.com', 'getnexhealth.com',
    'localmed.com', 'weave.com', 'revenuewell.com', 'solutionreach.com',
    'dentrixascend.com', 'opensquare.com', 'orthotrac.com',
    'carecredit.com/go',     'patientconnect365.com',
    'book.appointment-plus.com',
    'rula.com', 'cerebral.com',
    'talkspace.com', 'betterhelp.com', 'brightside.com',
    'neuroflow.com', 'valant.io', 'icanotes.com', 'therapynotes.com',
    'counselingsecure.com', 'therapynotes.com', 'ouremr.com',
    'drchrono.com', 'kidochealth.com', 'healthie.com', 'gethealthie.com',
    'practicebetter.io', 'janeapp.com', 'clinicense.com',
    'ptwired.com', 'webpt.com', 'prompt-emr.com', 'fusionwebclinic.com',
    'clinicient.com', 'raintreeinc.com', 'advancedmd.com',
    'modmed.com', 'ema.md', 'doxy.me',
    'patientnow.com', '4dhealthware.com',
    'onpatient.com', 'mdhq.com',
    # General appointment / calendar SaaS
    'calendly.com', 'acuityscheduling.com', 'squareup.com', 'square.site',
    'book.squareup.com', 'appointments.squareup.com', 'square.site/book',
    'simplybook.me', 'simplybook.it', 'setmore.com', 'jane.app',
    'schedulicity.com', 'appointlet.com', 'youcanbook.me', 'bookafy.com',
    'gettimely.com', 'timelyapp.com', 'bookingkoala.com',
    'microsoftbookings.com', 'outlook.office.com/book', 'bookings.microsoft.com',
    'twilio.com/appointments', 'oncehub.com', 'scheduleonce.com',
    'doodle.com', 'appointy.com', 'vcita.com', '10to8.com', 'picktime.com',
    'supersaas.com', 'appointmentplus.com', 'appointment-plus.com',
    'bookeo.com', 'reservio.com', 'octane.ai',
    'shopify.com/bookings', 'bookings.gettimely.com',
    'hubspot.com/meetings', 'meetings.hubspot.com',
    'chili-piper.com', 'savvycal.com', 'reclaim.ai',
    'calendly.com/d/', 'koalendar.com', 'bookingpage.com',
    'bookedin.com', 'appointfix.com', 'timepage.app',
    'when2meet.com', 'needto.io', 'tidycal.com',
    # Spa / salon / wellness / fitness
    'booker.com', 'go.booker.com', 'secure-booker.com', 'signin.booker.com',
    'mindbodyonline.com', 'mindbody.io', 'clients.mindbodyonline.com',
    'dashboard.mindbodyonline.com', 'booking.mindbodyonline.com',
    'fresha.com', 'booksy.com', 'vagaro.com', 'mangomint.com',
    'genbook.com', 'mytime.com', 'schedulista.com', 'glossgenius.com',
    'styleseat.com', 'shedul.com', 'treatwell.com',
    'treatwell.co.uk', 'planity.com', 'salonized.com', 'phorest.com',
    'shortcut.app', 'boulevard.io', 'joinblvd.com', 'zenoti.com',
    'daymart.com', 'spabooker.com', 'book4time.com', 'timekit.io',
    'meevo.com', 'rosysalonsoftware.com', 'shortcuts-software.com',
    'ondemand.mindbody.io', 'healcode.com', 'gymmaster.com',
    'gloomeme.com', 'punchpass.com', 'wellnessliving.com',
    'marianaatek.com', 'marianatek.com', 'zingfit.com', 'peakpass.com',
    'clubready.com', 'abcfinancial.com', 'mindbodyonline.com/classic',
    'schedulicity.com', 'genbook.com', 'booker.com/auth',
    'squareappointments.com', 'acuityscheduling.com/schedule',
    'setmore.com/book', 'timedeca.com', 'book.salonized.com',
    # Hotel / resort / dining booking engines
    'synxis.com', 'be.synxis.com',
    'travelclick.com', 'reservations.travelclick.com',
    'cloudbeds.com', 'us.cloudbeds.com', 'hotels.cloudbeds.com',
    'mews.com', 'app.mews.com', 'booking.mews.com',
    'sevenrooms.com', 'opentable.com', 'resy.com', 'tock.com',
    'book.bestwestern.com', 'reservations.ihotelier.com',
    'ihotelier.com', 'verticalbooking.com', 'simplebooking.it',
    'bookings.guestline.app', 'guestline.app', 'hotelrunner.com',
    'bookingsengine.com', 'reservhotel.com', 'roomkey.com',
    'peek.com', 'fareharbor.com', 'checkfront.com', 'bokun.io',
    'rezdy.com', 'regiondo.com', 'trekksoft.com', 'xola.com',
    # Australia / NZ healthcare
    'hotdoc.com.au', 'hotdoc.com', 'healthengine.com.au', 'healthengine.com',
    'cliniko.com', 'powerdiary.com', 'nookal.com', 'halaxy.com',
    'appointmind.com', 'automedsystems.com.au', 'medirecords.com',
    'pracsoft.com.au', 'genie-soft.com.au', 'bestpractice.com.au',
    # Europe / LATAM healthcare
    'doctolib.fr', 'doctolib.de', 'doctolib.com', 'doctolib.it',
    'jameda.de', 'samedi.de', 'terminland.de', 'doctena.com', 'doctena.lu',
    'dr-flex.de', 'arzt-termin.de', 'clickdoc.de', 'recasens.com',
    'doctoralia.com', 'doctoralia.es', 'doctoralia.com.mx', 'doctoralia.com.br',
    'mio-dottore.it', 'znanylekarz.pl', 'doktortakvimi.com', 'medicover.com',
    'topdoctors.es', 'topdoctors.co.uk', 'doctena.be', 'sanitas.es',
    'quirotel.com', 'cedigital.de', 'patientus.de',
    'doctorly.de', 'terminapp.de', 'online-termin.biz',
    # More US clinic / therapy / specialty schedulers
    'zipappointments.com', 'elationemr.com', 'app.elationemr.com',
    'patient.elationemr.com', 'intakeq.com', 'kirrahealth.com',
    'luminello.com', 'icanotes.com', 'therapyappointment.com',
    'secure.therapyappointment.com', 'clientsecure.me',
    'pockit.com', 'getpockit.com', 'bookapet.com',
    'appointment-plus.com', 'book.appointment-plus.com',
    'schedulicity.com/scheduling', 'genbook.com/bookings',
    'amzy.com', 'doctible.com', 'solutionreach.com',
    'demandforce.com', 'lumahealth.io', 'relatient.net',
    'kyruus.com', 'cedaron.com', 'sph.com',
    'epicmychart.com', 'opennotes.org',
    # GoHighLevel / CRM scheduling (common on med-spa & clinic sites)
    'leadconnectorhq.com', 'gohighlevel.com', 'msgsndr.com',
    'services.leadconnectorhq.com', 'api.leadconnectorhq.com',
    'widgets.leadconnectorhq.com', 'app.gohighlevel.com',
    # Cal.com, FlexBooker, Bookwhen, Timetap
    'cal.com', 'app.cal.com', 'flexbooker.com', 'bookwhen.com',
    'timetap.com', 'yocale.com', 'appointedd.com', 'simplymeet.me',
    'harmonizely.com', 'meetergo.com', 'meetfox.com',
    # Urgent care / queue / walk-in / clinic seat booking
    'solvhealth.com', 'solv.com', 'clockwisemd.com', 'waitwhile.com',
    'qless.com', 'qminder.com', 'skedulo.com', 'clockwise.md',
    'everseat.com', 'app.everseat.com',
    # WordPress / site-builder booking plugins (embedded widgets)
    'wpamelia.com', 'booking.wpamelia.com', 'bookly.pro', 'bookly.net',
    'latepoint.com', 'simplyschedule.me', 'bookingpressplugin.com',
    'ameliabooking.com', 'wpbookingcalendar.com', 'motopress.com',
    # Tebra, ModMed patient, Athena variants
    'tebra.com', 'patient.kareo.com', 'athenanet.athenahealth.com',
    'schedule.athenahealth.com', 'bookings.healow.com',
    # Therapy / behavioral health platforms
    'sondermind.com', 'lifestance.com', 'quenza.com', 'upvio.com',
    'practiceq.com', 'pabau.com', 'therapytribe.com', 'mentaltherapy.io',
    'breakthrough.com', 'regrouptherapy.com', 'inclusivehealth.com',
    # Fitness / studio / franchise
    'momence.com', 'glofox.com', 'arketa.co', 'walla.app', 'classpass.com',
    'punchpass.com', 'zingfit.com', 'marianatek.com', 'mindbody.io',
    'exercise.com', 'gymmasteronline.com',
    # Hospitality / tours
    'direct-book.com', 'littlehotelier.com', 'eviivo.com', 'roomraccoon.com',
    'staah.com', 'bookingjini.com', 'cultbooking.com', 'profitroom.com',
    'cbooking.com', 'bookingexpert.com', 'bookoloengine.com',
    # Dental / specialty
    'lighthouse360.com', 'dentrixascend.com', 'curvehero.com',
    'orthofi.com', 'patientviewer.com',
    # Pharmacy / retail clinic
    'cvs.com/minuteclinic', 'walgreens.com/findcare',
    'scheduling.cvs.com', 'appointment.walgreens.com',
    # Embeds / white-label schedulers
    'embed.ly', 'schedulista.com', 'booklikeaboss.com',
    '10to8.com', 'picktime.com', 'bookafy.com',
    'schedulock.com', 'zefi.ai', 'schedulr.com',
    'appointmentquest.com', 'appointmentcore.com',
    # Canada / UK common
    'jane.app', 'janeapp.com', 'cliniko.com', 'powerdiary.com',
    'nookal.com', 'halaxy.com', 'hotdoc.com.au',
]

# Login / patient-portal URLs are NOT online booking by themselves.
PORTAL_LOGIN_HINTS = (
    'logon', 'login', 'log-in', 'signin', 'sign-in', 'sign_in',
    '/account/log', '/auth/', 'sso.', 'forgot-password',
    'patientportal', 'patient-portal', 'patient_portal',
)

# Raw HTML/script markers for embedded booking widgets (checked before stripping scripts).
# Prefer concrete vendor domains over platform boilerplate.
BOOKING_WIDGET_MARKERS = [
    'clients.mindbodyonline.com', 'calendly.com', 'acuityscheduling.com',
    'hotdoc.com', 'healthengine.com', 'cliniko.com', 'powerdiary.com',
    'nookal.com', 'halaxy.com', 'fresha.com', 'jane.app',
    'doctolib.', 'jameda.de', 'samedi.de', 'terminland',
    'automedsystems.com.au', 'squarespace-scheduling',
    'go.booker.com', 'secure-booker.com', 'booker.com',
    'zenoti.com', 'boulevard.io', 'joinblvd.com', 'glossgenius.com',
    'styleseat.com', 'vagaro.com', 'mangomint.com', 'booksy.com',
    'healcode.com', 'mindbodyonline.com', 'vcita.com', 'simplybook.',
    'be.synxis.com', 'synxis.com', 'nexhealth.com', 'getnexhealth.com',
    'simplepractice.com', 'therapyappointment.com', 'theranest.com',
    'patientpop.com', 'healow.com', 'zocdoc.com', 'setmore.com',
    'schedulicity.com', 'appointy.com', 'bookeo.com', 'mews.com',
    'cloudbeds.com', 'fareharbor.com', 'opentable.com', 'sevenrooms.com',
    'healthie.com', 'gethealthie.com', 'practicebetter.io',
    'zipappointments.com', 'elationemr.com', 'app.elationemr.com',
    'clientsecure.me', 'therapyappointment.com',
    'cal.com', 'api.cal.com', 'app.cal.com', 'flexbooker.com',
    'leadconnectorhq.com', 'gohighlevel.com',
    'solvhealth.com', 'clockwisemd.com', 'waitwhile.com', 'wpamelia.com',
    'bookly.pro', 'latepoint.com', 'momence.com', 'yocale.com', 'timetap.com',
    'mentaltherapy.io', 'bookwhen.com', 'appointedd.com',
    # Clinic queue / waitlist / seat booking (seen on centennialmedical.com)
    'everseat.com', 'app.everseat.com',
]

# Wix (and similar builders) inject booking *platform* JS into every site even
# when no booking UI is shown on the frontend. Never treat these alone as YES.
WIX_BOOKING_BOILERPLATE = [
    'bookonlineviewerwidget',
    'bookings-service-list-widget',
    'bookings-service-details-widget',
    'bookingservicepageviewerwidget',
    'bookings.wixapps.net',
    'bookings-viewer/widget',
    'bookings-uou',
    'wix-bookings',
    'wixbookings',
]

# Forms that must never count as booking forms.
NON_BOOKING_FORM_HINTS = [
    'password', 'login', 'sign in', 'signin', 'sign-up', 'signup',
    'register', 'search', 'newsletter', 'subscribe', 'promo code',
    'coupon', 'cart', 'checkout', 'username',
]

PHONE_ONLY_MARKERS = [
    r'\bcall\b', r'\bphone\b', r'\bdial\b', r'\btel:', r'call us',
    r'please call', r'call to (?:book|schedule|make)',
    r'phone to (?:book|schedule)', r'ring us',
    # DE / FR / ES
    r'\banrufen\b', r'telefonisch', r'bitte anrufen',
    r'appelez', r'par téléphone', r'par telephone',
    r'llame', r'llamar al', r'por teléfono', r'por telefono',
]

ONLINE_CONTEXT_MARKERS = [
    'online', 'website', 'click', 'form', 'here', 'book now',
    'schedule now', 'available online', 'book online', 'schedule online',
    'jetzt', 'hier', 'online buchen', 'en ligne', 'en línea', 'en linea',
    'jetzt buchen', 'online termin', 'prenota', 'afspraak',
]

APPOINTMENT_FORM_FIELD_PATTERNS = {
    'date': [
        r'appointment[_ -]?date', r'visit[_ -]?date', r'appt[_ -]?date',
        r'pref(?:erred)?[_ -]?date', r'requested[_ -]?date', r'schedule[_ -]?date',
        r'\bdate\b',
    ],
    'time': [
        r'appointment[_ -]?time', r'visit[_ -]?time', r'appt[_ -]?time',
        r'pref(?:erred)?[_ -]?time', r'requested[_ -]?time', r'schedule[_ -]?time',
        r'time[_ -]?slot', r'timeslot', r'\btime\b', r'\bslot\b',
    ],
    'provider': [
        r'provider', r'doctor', r'physician', r'specialist', r'clinician',
        r'therapist', r'practitioner', r'dentist', r'consultant',
        r'care[_ -]?team', r'provider[_ -]?name', r'physician[_ -]?name',
        r'doctor[_ -]?name',
    ],
    'service': [
        r'service', r'appointment[_ -]?type', r'visit[_ -]?type', r'appt[_ -]?type',
        r'type[_ -]?of[_ -]?visit', r'procedure', r'treatment',
        r'reason[_ -]?for[_ -]?visit', r'visit[_ -]?reason', r'department',
        r'service[_ -]?type', r'\breason\b',
    ],
}

CONTACT_FORM_ONLY_INDICATORS = [
    'name', 'email', 'phone', 'message', 'comment', 'subject', 'inquiry',
]


class Result(Enum):
    YES = "YES"
    NO = "NO"
    ERROR = "ERROR"
    SKIP = "SKIP"


@dataclass
class WebsiteResult:
    url: str
    result: Result
    reason: str = ""


class GracefulShutdown:
    def __init__(self):
        self.shutdown_requested = False

    def request_shutdown(self, signum, frame):
        logger.info("Shutdown requested. Finishing current operation...")
        self.shutdown_requested = True


shutdown_handler = GracefulShutdown()
signal.signal(signal.SIGINT, shutdown_handler.request_shutdown)
signal.signal(signal.SIGTERM, shutdown_handler.request_shutdown)


def _force_ipv4() -> None:
    """Prefer IPv4 — some hosts fail with 'Network is unreachable' on IPv6."""
    import socket
    if getattr(socket, '_scraper_ipv4_patched', False):
        return
    _orig = socket.getaddrinfo

    def getaddrinfo_ipv4(host, port, family=0, type=0, proto=0, flags=0):
        return _orig(host, port, socket.AF_INET, type, proto, flags)

    socket.getaddrinfo = getaddrinfo_ipv4
    socket._scraper_ipv4_patched = True


_force_ipv4()


class WebsiteFetcher:
    def __init__(
        self,
        timeout: int = REQUEST_TIMEOUT,
        max_retries: int = MAX_NETWORK_RETRIES,
    ):
        self.timeout = timeout
        self.max_retries = max_retries
        self.session = requests.Session()
        self.session.headers.update({
            'User-Agent': (
                'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                'AppleWebKit/537.36 (KHTML, like Gecko) '
                'Chrome/120.0.0.0 Safari/537.36'
            ),
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8',
            'Accept-Language': 'en-US,en;q=0.9,de;q=0.8,fr;q=0.7,es;q=0.6,it;q=0.5',
        })
        self._curl = None
        self._alt_cache: dict = {}
        try:
            from curl_cffi import requests as curl_requests
            self._curl = curl_requests
        except Exception:
            self._curl = None

    @staticmethod
    def is_challenge_page(html: Optional[str]) -> bool:
        """True for captcha / bot-wall / forbidden shells (not a real site).

        Large pages that merely load Cloudflare scripts must NOT be flagged —
        that caused mass false ERRORs (finnegancounseling, mindworkstx, etc.).

        Hard 403 Forbidden shells (e.g. SiteGround after UC) MUST be flagged —
        scoring them as NO caused false negatives like bfsuccess.com.
        """
        if not html:
            return True
        low = html.lower()
        n = len(html)
        if n < 800:
            return True

        title = ''
        if '<title' in low:
            m = re.search(r'<title[^>]*>([^<]{0,120})</title>', low, re.I)
            if m:
                title = m.group(1).strip()
        title_norm = re.sub(r'[^a-z0-9]+', ' ', title.lower()).strip()

        # Hard block pages — any size (UC often returns a styled 403 shell).
        forbidden_titles = (
            '403 forbidden', '403', 'access denied', 'forbidden',
            'robot challenge screen', 'robot check', 'security check',
            'just a moment', 'attention required',
        )
        if title_norm in forbidden_titles or title_norm.startswith('403'):
            return True
        if 'access to this page is forbidden' in low:
            return True
        if 'sg-captcha' in low or '/.well-known/sgcaptcha' in low:
            return True

        interstitial_titles = (
            'just a moment...',
            'attention required',
            'access denied',
            '403 forbidden',
            'robot check',
            'security check',
        )
        if n < MIN_REAL_PAGE_BYTES and (
            title.lower() in interstitial_titles
            or title.lower().startswith('attention required')
        ):
            return True

        # Strong phrases — still ignore on large pages unless title is interstitial.
        if any(m in low for m in CHALLENGE_MARKERS_STRONG):
            if n < MIN_REAL_PAGE_BYTES or title.lower() in interstitial_titles:
                return True

        # Weak CF script markers: only on small pages.
        if n < 8000 and any(m in low for m in CHALLENGE_MARKERS_WEAK):
            return True

        # SiteGround / similar soft-block: meta refresh into captcha.
        if 'http-equiv="refresh"' in low and (
            'captcha' in low or 'challenge' in low or 'sgcaptcha' in low
        ):
            if n < 8000:
                return True
        return False

    def normalize_url(self, url: str) -> Optional[str]:
        url = (url or '').strip()
        if not url or url.lower() in ('n/a', 'na', 'none', '-', 'null'):
            return None

        # Obvious non-website values sometimes appear in export sheets.
        low = url.lower()
        if any(bad in low for bad in (
            'linkedin.com/', 'mail.google.com/', 'mailto:', 'javascript:',
        )):
            return None

        if not url.startswith(('http://', 'https://')):
            url = 'https://' + url

        url = url.rstrip('/')

        try:
            parsed = urlparse(url)
            if not parsed.netloc or '.' not in parsed.netloc:
                return None
            return url
        except Exception:
            return None

    def same_site(self, url_a: str, url_b: str) -> bool:
        try:
            a = _strip_www(urlparse(url_a).netloc)
            b = _strip_www(urlparse(url_b).netloc)
            return bool(a) and a == b
        except Exception:
            return False

    def _candidate_urls(self, url: str) -> List[str]:
        """Build URL variants: https/http and www/non-www."""
        normalized = self.normalize_url(url)
        if not normalized:
            return []

        parsed = urlparse(normalized)
        host = parsed.netloc
        path = parsed.path if parsed.path not in ('', '/') else ''
        if path.endswith('/'):
            path = path.rstrip('/')
        query = f'?{parsed.query}' if parsed.query else ''
        suffix = f'{path}{query}'

        bare = host[4:] if host.startswith('www.') else host
        hosts = [bare, f'www.{bare}']

        candidates: List[str] = []
        for scheme in ('https', 'http'):
            for h in hosts:
                candidates.append(f'{scheme}://{h}{suffix}')

        seen: Set[str] = set()
        ordered: List[str] = []
        for c in candidates:
            key = c.lower()
            if key not in seen:
                seen.add(key)
                ordered.append(c)
        return ordered

    _IMPERSONATE_ROTATION = (
        'chrome131', 'chrome124', 'chrome120', 'chrome136', 'safari17_0',
    )

    def _get_once(
        self,
        url: str,
        verify: bool = True,
        impersonate: Optional[str] = None,
        read_timeout: Optional[float] = None,
    ):
        # Prefer Chrome TLS fingerprint when available (bypasses some bot walls).
        timeout = (CONNECT_TIMEOUT, read_timeout or self.timeout)
        # Broad Accept avoids intermittent 406 Not Acceptable on picky CDNs.
        # Avoid 'br' — some curl builds hang decompressing brotli.
        headers = {
            'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,'
                      'image/avif,image/webp,image/apng,*/*;q=0.8',
            'Accept-Language': 'en-US,en;q=0.9',
            'Accept-Encoding': 'gzip, deflate',
            'Cache-Control': 'no-cache',
            'Pragma': 'no-cache',
            'Upgrade-Insecure-Requests': '1',
            'Sec-Fetch-Dest': 'document',
            'Sec-Fetch-Mode': 'navigate',
            'Sec-Fetch-Site': 'none',
            'Sec-Fetch-User': '?1',
        }
        if self._curl is not None:
            try:
                return self._curl.get(
                    url,
                    timeout=timeout,
                    allow_redirects=True,
                    verify=verify,
                    impersonate=impersonate or 'chrome131',
                    headers=headers,
                )
            except Exception:
                pass
        return self.session.get(
            url, timeout=timeout, allow_redirects=True, verify=verify, headers=headers
        )

    @staticmethod
    def _backoff(attempt: int) -> None:
        time.sleep(min(2.0, FLAKY_FETCH_BACKOFF * (attempt + 1)))

    def discover_alternate_urls(self, url: str, limit: int = 3) -> List[str]:
        """When a domain is blocked/unreachable, find likely live pages via search.

        Many spa chains use a vanity domain that redirects in browsers but is
        unreachable from datacenter IPs. Search often surfaces the real brand URL.
        """
        if not ALTERNATE_DISCOVERY:
            return []

        normalized = self.normalize_url(url)
        if not normalized:
            return []

        host = _strip_www(urlparse(normalized).netloc)
        cache_key = f'{host}|{limit}'
        if cache_key in self._alt_cache:
            return list(self._alt_cache[cache_key])

        base_label = host.split('.')[0]
        # Two focused queries are enough; Booker/Mindbody show up under "book".
        queries = [
            host,
            f'{base_label} book appointment',
        ]

        from urllib.parse import parse_qs, unquote

        booking_domain_needles = tuple(
            d.lower() for d in THIRD_PARTY_BOOKING_DOMAINS if '.' in d
        )

        def shares_brand(label: str, host_path: str) -> bool:
            """True when alternate URL clearly belongs to the same brand.

            Full label match, or vanity-domain stem after stripping a known
            location suffix (4everyoungchandler → 4everyoung…). Never use
            short mid-string prefixes — those matched
            pregnancyresourcecenter → pregnancyresourcekeene (false YES).
            """
            label = re.sub(r'[^a-z0-9]', '', label.lower())
            host_alnum = re.sub(r'[^a-z0-9]', '', host_path.lower())
            if not label or len(label) < 4:
                return False
            if label in host_alnum:
                return True
            # Strip trailing city/region tokens used on vanity med-spa domains.
            location_suffixes = (
                'chandler', 'dallas', 'houston', 'austin', 'phoenix', 'scottsdale',
                'miami', 'tampa', 'orlando', 'denver', 'seattle', 'portland',
                'chicago', 'boston', 'atlanta', 'nashville', 'charlotte',
                'pavilions', 'downtown', 'midtown',
            )
            stem = label
            for suf in location_suffixes:
                if label.endswith(suf) and len(label) - len(suf) >= 6:
                    stem = label[: -len(suf)]
                    break
            if stem != label and len(stem) >= 6 and stem in host_alnum:
                return True
            return False

        def unwrap_href(href: str) -> str:
            href = (href or '').strip()
            if href.startswith('//'):
                href = 'https:' + href
            # Decode DDG redirect wrappers (/l/?uddg=...)
            if 'uddg=' in href:
                qs = parse_qs(urlparse(href).query)
                if 'uddg' in qs:
                    href = unquote(qs['uddg'][0])
            return href

        def is_booking_platform(h: str, path: str) -> bool:
            h = _strip_www(h or '')
            path = (path or '').lower()
            for d in booking_domain_needles:
                if '/' in d:
                    root, _, rest = d.partition('/')
                    if (h == root or h.endswith('.' + root)) and rest in path:
                        return True
                    continue
                if h == d or h.endswith('.' + d):
                    return True
            return False

        def rank(href: str) -> tuple:
            """Higher score first: real booking platforms beat staging/marketing."""
            p = urlparse(href)
            h = _strip_www(p.netloc)
            path = (p.path or '').lower()
            score = 0
            if is_booking_platform(h, path):
                score += 100
            if any(k in path for k in ('book', 'appoint', 'schedule', 'reserv')):
                score += 20
            if h.startswith('staging.') or h.startswith('preview.') or 'staging' in h:
                score -= 50
            if any(x in h for x in ('facebook.', 'instagram.', 'linkedin.')):
                score -= 100
            return (-score, href)

        scored: List[str] = []
        seen: Set[str] = set()

        def add_candidate(href: str) -> None:
            href = (href or '').rstrip('/')
            if not href:
                return
            key = href.lower()
            if key in seen:
                return
            seen.add(key)
            scored.append(href)
            # Booker marketing sites often link secure-booker.com/{slug}/... which
            # is unreachable here; go.booker.com/location/{slug} is the live UI
            # (only derive when search already proved the slug exists on Booker).
            try:
                p = urlparse(href)
                h = _strip_www(p.netloc)
                parts = [x for x in (p.path or '').split('/') if x]
                if 'secure-booker.com' in h and parts:
                    add_candidate(f'https://go.booker.com/location/{parts[0]}')
                if h.endswith('booker.com') and 'location' in parts:
                    idx = parts.index('location')
                    if idx + 1 < len(parts):
                        add_candidate(f'https://go.booker.com/location/{parts[idx + 1]}')
            except Exception:
                pass

        for query in queries:
            try:
                if self._curl is not None:
                    resp = self._curl.get(
                        'https://html.duckduckgo.com/html/',
                        params={'q': query},
                        timeout=min(20, self.timeout + 5),
                        impersonate='chrome120',
                    )
                    html = resp.text
                else:
                    resp = self.session.get(
                        'https://html.duckduckgo.com/html/',
                        params={'q': query},
                        timeout=min(20, self.timeout + 5),
                    )
                    html = resp.text
            except Exception:
                continue

            soup = BeautifulSoup(html, 'lxml')
            for a in soup.find_all('a', href=True):
                href = unwrap_href(a['href'])
                if not href.startswith('http'):
                    continue
                parsed = urlparse(href)
                h = _strip_www(parsed.netloc)
                if not h or h.endswith('duckduckgo.com') or h.endswith('duck.com'):
                    continue
                if h == host:
                    continue
                if any(x in h for x in (
                    'facebook.com', 'instagram.com', 'linkedin.com',
                    'youtube.com', 'twitter.com', 'x.com', 'yelp.com',
                    'wheree.com', 'chamber.com', 'skincareandspa.com',
                    'atly.com', 'mapquest.com', 'yellowpages.com',
                    'bbb.org', 'wellness.com', 'play.google.com',
                    'campusbooker.com', 'tradecolleges.org', 'skozo.ai',
                    'findglocal.com', 'collegefitproject.com', 'solo.to',
                    'dougcosocial.com', 'updowntoday.com', 'allbiz.com',
                    'yellowpages.com', 'bbb.org', 'mapquest.com',
                    'prospeo.io', 'linktr.ee', 'francentral.com',
                    'americasbestfranchises.com', 'franchisetop40.com',
                )):
                    continue
                # Booking-platform hits only need brand in host/path;
                # other alts must share brand too (already required below).
                if not shares_brand(base_label, h + parsed.path):
                    continue
                add_candidate(href)

            # One real scheduler URL is enough — skip remaining search queries.
            if any(is_booking_platform(urlparse(u).netloc, urlparse(u).path) for u in scored):
                break

        scored.sort(key=rank)
        result = scored[:limit]
        self._alt_cache[cache_key] = list(result)
        return result

    @staticmethod
    def _trim_html(html: str) -> str:
        """Keep enough HTML for detection without holding multi‑MB pages in RAM."""
        if html and len(html) > MAX_HTML_BYTES:
            return html[:MAX_HTML_BYTES]
        return html

    def _handle_response(self, response) -> Tuple[Optional[str], Optional[str], Optional[str]]:
        """Normalize a successful HTTP response into (html, url, error).

        curl_cffi raises its own HTTPError (not requests.HTTPError), so callers
        should prefer checking status_code here instead of raise_for_status alone.
        Large non-challenge bodies are accepted even on 403/406 — some hosts
        gate bots with status codes but still return usable HTML after retries.
        """
        status = getattr(response, 'status_code', None)
        html = getattr(response, 'text', '') or ''
        final = getattr(response, 'url', None)
        challenge = self.is_challenge_page(html)
        headers = {str(k).lower(): str(v).lower() for k, v in (getattr(response, 'headers', {}) or {}).items()}
        sg_challenge = headers.get('sg-captcha') == 'challenge'

        # SiteGround soft-block often returns HTTP 202 + tiny meta-refresh.
        if sg_challenge or (status == 202 and (challenge or len(html) < 2000)):
            return None, None, "Bot protection / captcha challenge"

        if status and status >= 400:
            # Real page content slipped through a soft block — use it.
            if html and not challenge and len(html) >= MIN_REAL_PAGE_BYTES:
                return self._trim_html(html), final, None
            if status in (401, 403, 429) and (challenge or len(html) < MIN_REAL_PAGE_BYTES):
                return None, None, "Bot protection / captcha challenge"
            return None, None, f"HTTP {status}"

        if challenge:
            return None, None, "Bot protection / captcha challenge"
        return self._trim_html(html), final, None

    def fetch(self, url: str) -> Tuple[Optional[str], Optional[str], Optional[str]]:
        """Return (html, final_url, error). Tries https/www/http variants.

        Flaky hosts (connection drops, intermittent 403/406) are retried across
        several rounds with rotating TLS fingerprints — mirrors refreshing a
        browser several times before giving up.
        """
        candidates = self._candidate_urls(url)
        if not candidates:
            return None, None, "Invalid URL format"

        last_error = "Max retries exceeded"
        saw_challenge = False
        retryable_status = {401, 403, 406, 408, 429, 500, 502, 503, 504}
        dns_fails = 0
        saw_retryable = False
        deadline = time.monotonic() + MAX_FETCH_SECONDS

        for round_idx in range(FLAKY_FETCH_ROUNDS):
            if time.monotonic() >= deadline:
                break
            impersonate = self._IMPERSONATE_ROTATION[
                round_idx % len(self._IMPERSONATE_ROTATION)
            ]
            hard_connect_fails = 0
            round_retryable = False
            # Later rounds: https only (skip http duplicates) + shorter timeout.
            if round_idx == 0:
                round_candidates = candidates
            else:
                round_candidates = [c for c in candidates if c.startswith('https://')] or candidates
            read_to = FLAKY_READ_TIMEOUT if round_idx > 0 else self.timeout

            for candidate in round_candidates:
                if time.monotonic() >= deadline:
                    last_error = last_error if last_error != "Max retries exceeded" else "Request timed out"
                    break
                try:
                    response = self._get_once(
                        candidate,
                        verify=True,
                        impersonate=impersonate,
                        read_timeout=read_to,
                    )
                    html, final, err = self._handle_response(response)
                    if err:
                        last_error = err
                        low = err.lower()
                        if 'captcha' in low or 'bot protection' in low:
                            saw_challenge = True
                            round_retryable = True
                        elif err.startswith('HTTP '):
                            try:
                                code = int(err.split()[1])
                            except (IndexError, ValueError):
                                code = 0
                            if code in retryable_status:
                                round_retryable = True
                            # 404/410 etc. — try next URL variant, not more rounds.
                        continue
                    return html, final, None
                except requests.exceptions.Timeout:
                    last_error = "Request timed out"
                    hard_connect_fails += 1
                    round_retryable = True
                    if hard_connect_fails >= 2:
                        break  # host is slow/dead — next round or give up
                    continue
                except requests.exceptions.ConnectionError:
                    last_error = "Connection error"
                    hard_connect_fails += 1
                    round_retryable = True
                    if hard_connect_fails >= 3:
                        break
                    continue
                except requests.exceptions.SSLError:
                    try:
                        response = self._get_once(
                            candidate,
                            verify=False,
                            impersonate=impersonate,
                            read_timeout=read_to,
                        )
                        html, final, err = self._handle_response(response)
                        if err:
                            last_error = err
                            if 'captcha' in err.lower() or 'bot protection' in err.lower():
                                saw_challenge = True
                                round_retryable = True
                            continue
                        return html, final, None
                    except Exception as e:
                        last_error = f"SSL error: {type(e).__name__}"
                        hard_connect_fails += 1
                        round_retryable = True
                        continue
                except Exception as e:
                    # curl_cffi DNSError / HTTPError / CertificateVerifyError, etc.
                    resp = getattr(e, 'response', None)
                    status = getattr(resp, 'status_code', None) if resp is not None else None
                    name = type(e).__name__
                    msg = str(e).lower()
                    if status:
                        if status in (401, 403, 429) and (
                            not resp or len(getattr(resp, 'text', '') or '') < MIN_REAL_PAGE_BYTES
                        ):
                            last_error = "Bot protection / captcha challenge"
                            saw_challenge = True
                            round_retryable = True
                        else:
                            last_error = f"HTTP {status}"
                            if status in retryable_status:
                                round_retryable = True
                    elif 'Timeout' in name or 'timed out' in msg:
                        last_error = "Request timed out"
                        hard_connect_fails += 1
                        round_retryable = True
                        if hard_connect_fails >= 2:
                            break
                    elif 'DNS' in name or 'Resolve' in name:
                        last_error = "Connection error"
                        hard_connect_fails += 1
                        dns_fails += 1
                    elif 'Certificate' in name or 'SSL' in name:
                        try:
                            response = self._get_once(
                                candidate,
                                verify=False,
                                impersonate=impersonate,
                                read_timeout=read_to,
                            )
                            html, final, err = self._handle_response(response)
                            if not err:
                                return html, final, None
                            last_error = err or f"SSL error: {name}"
                            round_retryable = True
                        except Exception:
                            last_error = f"SSL error: {name}"
                            hard_connect_fails += 1
                            round_retryable = True
                    else:
                        last_error = f"Request failed: {name}"
                        hard_connect_fails += 1
                        round_retryable = True
                    continue

            saw_retryable = saw_retryable or round_retryable
            # Dead DNS across every variant — further rounds won't help.
            if dns_fails >= len(candidates):
                break
            # Only soft failures (conn/403/challenge/5xx) deserve another refresh.
            if not round_retryable:
                break
            if round_idx < FLAKY_FETCH_ROUNDS - 1:
                self._backoff(round_idx)

        if saw_challenge:
            return None, None, "Bot protection / captcha challenge"
        return None, None, last_error


class BookingDetector:
    def __init__(self):
        self.booking_keywords_lower = [k.lower() for k in BOOKING_KEYWORDS]
        self.strong_ctas = [k.lower() for k in STRONG_ONLINE_CTAS]
        self.strong_labels = {k.lower() for k in STRONG_LINK_LABELS}
        self.booking_url_patterns = [re.compile(p, re.IGNORECASE) for p in BOOKING_URL_PATTERNS]
        self.booking_host_patterns = [re.compile(p, re.IGNORECASE) for p in BOOKING_HOST_PATTERNS]
        self.follow_patterns = [re.compile(p, re.IGNORECASE) for p in FOLLOW_LINK_PATTERNS]
        # Host-boundary patterns so campusbooker.com does NOT match booker.com.
        self.third_party_patterns = []
        for d in THIRD_PARTY_BOOKING_DOMAINS:
            d = d.lower().strip()
            if not d:
                continue
            if '/' in d:
                self.third_party_patterns.append(
                    re.compile(re.escape(d), re.IGNORECASE)
                )
            else:
                self.third_party_patterns.append(
                    re.compile(
                        r'(?<![a-z0-9-])' + re.escape(d) + r'(?![a-z0-9-])',
                        re.IGNORECASE,
                    )
                )
        self.widget_markers = [m.lower() for m in BOOKING_WIDGET_MARKERS]
        # Boundary-aware patterns so "cal.com" does not match "*medical.com".
        self.widget_marker_patterns = []
        for m in self.widget_markers:
            if '.' in m and '/' not in m:
                self.widget_marker_patterns.append(
                    (m, re.compile(
                        r'(?<![a-z0-9-])' + re.escape(m) + r'(?![a-z0-9-])',
                        re.IGNORECASE,
                    ))
                )
            else:
                self.widget_marker_patterns.append((m, None))
        self.phone_only_re = [re.compile(p, re.IGNORECASE) for p in PHONE_ONLY_MARKERS]

    def detect(self, html: str, base_url: str) -> Tuple[Result, str]:
        if not html:
            return Result.ERROR, "No HTML content received"

        soup = BeautifulSoup(html, 'lxml')

        # Check embeds / scripts / raw widget markers before stripping them.
        iframe_hit = self._has_booking_iframe(soup)
        if iframe_hit:
            return Result.YES, iframe_hit

        script_hit = self._has_booking_script(soup)
        if script_hit:
            return Result.YES, script_hit

        widget_hit = self._has_booking_widget_marker(html)
        if widget_hit:
            return Result.YES, widget_hit

        for tag in soup.find_all(['script', 'style', 'noscript']):
            tag.decompose()

        text_content = soup.get_text(separator=' ', strip=True).lower()
        links = soup.find_all('a', href=True)
        buttons = soup.find_all(['button', 'input'])
        forms = soup.find_all('form')

        if self._has_third_party_booking(links):
            return Result.YES, "Found third-party booking link"

        if self._has_booking_link(links):
            return Result.YES, "Found booking-related link"

        if self._has_booking_button(buttons):
            return Result.YES, "Found booking-related button"

        if self._has_booking_form(forms):
            return Result.YES, "Found appointment booking form"

        # Body-copy-only matches are too noisy for YES (CMS leftovers, policies).
        # Keep as soft signal for Playwright fallback, not a direct YES.
        if self._has_contact_form_only(forms, text_content):
            return Result.NO, "Only contact form found (not booking)"

        return Result.NO, "No booking feature detected"

    def looks_js_heavy(self, html: str) -> bool:
        """True when the page is likely an SPA / JS-rendered shell."""
        if not html:
            return False
        soup = BeautifulSoup(html, 'lxml')
        scripts = soup.find_all('script')
        for tag in soup.find_all(['script', 'style', 'noscript']):
            tag.decompose()
        text = soup.get_text(' ', strip=True)
        if len(html) > 2000 and len(text) < 400 and len(scripts) >= 5:
            return True
        root = soup.find(id=re.compile(r'^(root|app|__next|__nuxt)$', re.I))
        if root and len(text) < 600:
            return True
        return False

    def has_soft_booking_signal(self, html: str) -> bool:
        """Strong hints that justify a JS render (visible text / vendor only)."""
        if not html:
            return False

        # Vendor domains in raw HTML are fine (real embeds).
        low = html.lower()
        vendors = [
            'clients.mindbodyonline.com', 'calendly.com', 'acuityscheduling',
            'hotdoc.com', 'healthengine.com', 'cliniko.com', 'powerdiary.com',
            'nookal.com', 'halaxy.com', 'doctolib.', 'squarespace-scheduling',
            'data-booking', 'go.booker.com', 'secure-booker.com', 'zenoti.com',
            'boulevard.io', 'fresha.com', 'vagaro.com', 'jane.app',
            'leadconnectorhq.com', 'gohighlevel.com', 'cal.com', 'flexbooker.com',
            'solvhealth.com', 'wpamelia.com', 'bookly', 'latepoint',
        ]
        if any(v in low for v in vendors):
            return True

        # CTA phrases must appear in *visible* text — Wix JSON often contains
        # "Book Online" in platform config with no frontend booking UI.
        soup = BeautifulSoup(html, 'lxml')
        for tag in soup.find_all(['script', 'style', 'noscript']):
            tag.decompose()
        visible = soup.get_text(' ', strip=True).lower()
        phrases = [
            'book now', 'book online', 'schedule online', 'online booking',
            'book an appointment', 'book appointment', 'schedule an appointment',
            'request an appointment', 'termin buchen', 'prendre rendez-vous',
            'make a reservation', 'book your stay', 'check availability',
            'book a room', 'book a treatment', 'reserve now',
            'virtual visit', 'telehealth appointment', 'schedule a tour',
            'start therapy', 'get care now', 'online intake',
        ]
        return any(p in visible for p in phrases)

    def collect_follow_urls(self, html: str, base_url: str, limit: int = MAX_PAGES_PER_SITE - 1) -> List[str]:
        """Pick a few high-value internal pages to inspect next."""
        if not html or limit <= 0:
            return []

        soup = BeautifulSoup(html, 'lxml')
        candidates: List[Tuple[int, str]] = []
        seen: Set[str] = set()
        base_norm = base_url.rstrip('/')

        for link in soup.find_all('a', href=True):
            href = link['href'].strip()
            if not href or href.startswith(('#', 'mailto:', 'tel:', 'javascript:')):
                continue

            absolute = urljoin(base_url, href)
            absolute, _ = urldefrag(absolute)
            absolute = absolute.rstrip('/')

            if absolute.lower() == base_norm.lower():
                continue
            if absolute in seen:
                continue

            parsed = urlparse(absolute)
            if not parsed.scheme.startswith('http'):
                continue

            link_text = link.get_text(strip=True).lower()
            haystack = f'{link_text} {href}'.lower()

            score = 0
            for kw in self.strong_ctas:
                if kw in haystack:
                    score += 5
            for pattern in self.booking_url_patterns:
                if pattern.search(href):
                    score += 4
            for pattern in self.follow_patterns:
                if pattern.search(haystack):
                    score += 2

            if score <= 0:
                continue

            base_host = _strip_www(urlparse(base_url).netloc)
            link_host = _strip_www(parsed.netloc)
            if base_host and link_host and base_host != link_host:
                continue

            seen.add(absolute)
            candidates.append((score, absolute))

        candidates.sort(key=lambda x: (-x[0], x[1]))
        return [url for _, url in candidates[:limit]]

    def _matches_third_party(self, url: str) -> bool:
        return any(p.search(url) for p in self.third_party_patterns)

    @staticmethod
    def _is_portal_login_url(href: str, link_text: str = '') -> bool:
        """True for patient-portal / login links (not online booking CTAs)."""
        blob = f'{href or ""} {link_text or ""}'.lower()
        if not blob.strip():
            return False
        # Explicit booking/scheduling intent wins even on portal domains.
        if re.search(
            r'appoint|schedule|booking|\bbook\b|reserv|request.?visit',
            blob,
        ):
            return False
        return any(h in blob for h in PORTAL_LOGIN_HINTS)

    def _is_booking_host(self, url: str) -> bool:
        try:
            host = urlparse(url).netloc.lower()
        except Exception:
            return False
        if not host:
            return False
        return any(p.search(host) for p in self.booking_host_patterns)

    def _has_booking_iframe(self, soup: BeautifulSoup) -> Optional[str]:
        for iframe in soup.find_all('iframe'):
            src = iframe.get('src') or iframe.get('data-src') or ''
            if not src:
                continue
            src_l = src.lower()
            # Real embedded Wix Bookings iframe on the page (not platform JS config).
            if 'bookings.wixapps.net' in src_l or 'wix-bookings' in src_l:
                return f"Found booking iframe ({urlparse(src).netloc or 'embed'})"
            if self._matches_third_party(src) or self._is_booking_host(src):
                return f"Found booking iframe ({urlparse(src).netloc or 'embed'})"
            if any(p.search(src_l) for p in self.booking_url_patterns):
                return "Found booking iframe"
        return None

    def _has_booking_script(self, soup: BeautifulSoup) -> Optional[str]:
        for script in soup.find_all('script'):
            src = script.get('src') or ''
            content = script.string or ''
            blob = f'{src} {content}'.lower()

            # Ignore Wix platform booking registry noise.
            if any(noise in blob for noise in WIX_BOOKING_BOILERPLATE):
                # Still allow a hit if a *non-Wix* vendor domain is also present.
                cleaned = blob
                for noise in WIX_BOOKING_BOILERPLATE:
                    cleaned = cleaned.replace(noise, ' ')
                if self._matches_third_party(cleaned):
                    return "Found third-party booking script/widget"
                if self._widget_marker_in(cleaned):
                    return "Found embedded booking widget"
                continue

            if self._matches_third_party(blob):
                return "Found third-party booking script/widget"
            if self._widget_marker_in(blob):
                return "Found embedded booking widget"
        return None

    def _widget_marker_in(self, text: str) -> Optional[str]:
        """Return matched marker if present (host-boundary safe for domains)."""
        if not text:
            return None
        for marker, pattern in self.widget_marker_patterns:
            if pattern is not None:
                if pattern.search(text):
                    return marker
            elif marker in text:
                return marker
        return None

    def _has_booking_widget_marker(self, html: str) -> Optional[str]:
        html_l = html.lower()
        # Strip known Wix boilerplate spans so leftover tokens cannot match.
        cleaned = html_l
        for noise in WIX_BOOKING_BOILERPLATE:
            cleaned = cleaned.replace(noise, ' ')

        hit = self._widget_marker_in(cleaned)
        if hit:
            return f"Found booking widget marker ({hit})"
        return None

    @staticmethod
    def _is_contact_only_href(href: str) -> bool:
        """True for /contact or /contact-us style paths (not a scheduler URL)."""
        if not href:
            return False
        path = urlparse(href).path if '://' in href else href
        path = (path or '').rstrip('/').lower()
        # Hash-only anchors on contact pages still count as contact.
        if path.startswith('#'):
            return False
        return bool(re.search(
            r'(?:^|/)(?:contact|contact-us|contactus|get-in-touch|getintouch)$',
            path,
        ))

    def _has_booking_link(self, links: List) -> bool:
        for link in links:
            href = link.get('href', '') or ''
            link_text = link.get_text(strip=True).lower()
            label = re.sub(r'\s+', ' ', link_text).strip(' !.?:')

            # Ignore bookmark / social noise
            if 'bookmark' in label or 'facebook' in href or 'instagram' in href:
                continue
            # Patient portal / login alone is not online booking.
            if self._is_portal_login_url(href, link_text):
                continue

            if self._matches_third_party(href) or self._is_booking_host(href):
                return True

            for pattern in self.booking_url_patterns:
                if pattern.search(href):
                    return True

            # "Make an appointment" → /contact is a contact form, not booking
            # (aamirahomecare.com, fairwindscenter.org).
            contact_only = self._is_contact_only_href(href)

            if label in self.strong_labels and not contact_only:
                return True

            for keyword in self.strong_ctas:
                if keyword in link_text and not contact_only:
                    return True

            for keyword in self.booking_keywords_lower:
                if (
                    keyword in link_text
                    and not self._looks_phone_only(link_text)
                    and not contact_only
                ):
                    return True

        return False

    def _has_booking_button(self, buttons: List) -> bool:
        for button in buttons:
            if button.name == 'input':
                itype = (button.get('type') or '').lower()
                if itype in ('hidden', 'checkbox', 'radio', 'password', 'email', 'tel'):
                    continue
                button_text = (
                    (button.get('value') or '') + ' ' +
                    (button.get('aria-label') or '') + ' ' +
                    (button.get('name') or '')
                ).lower()
            else:
                button_text = button.get_text(strip=True).lower()
                button_text += ' ' + (button.get('aria-label') or '').lower()

            if 'bookmark' in button_text:
                continue

            for keyword in self.strong_ctas:
                if keyword in button_text:
                    return True

            for keyword in self.booking_keywords_lower:
                if keyword in button_text and not self._looks_phone_only(button_text):
                    return True

            # Require word-boundary book/schedule (avoids "Bookmarks").
            if (
                re.search(r'\b(book|schedule|booking|termin|cita|prenota)\b', button_text)
                and 'contact' not in button_text
                and not self._looks_phone_only(button_text)
            ):
                return True

        return False

    def _looks_phone_only(self, text: str) -> bool:
        text_l = text.lower()
        has_phone = any(p.search(text_l) for p in self.phone_only_re)
        if not has_phone:
            return False
        has_online = any(m in text_l for m in ONLINE_CONTEXT_MARKERS)
        return not has_online

    def _has_third_party_booking(self, links: List) -> bool:
        for link in links:
            href = link.get('href', '') or ''
            if not (self._matches_third_party(href) or self._is_booking_host(href)):
                continue

            link_text = link.get_text(strip=True).lower()
            # Skip EHR patient-portal login pages (AdvancedMD LogOn, MyChart login, etc.)
            if self._is_portal_login_url(href, link_text):
                continue

            href_l = href.lower()
            if any(word in link_text for word in (
                'book', 'appointment', 'schedule', 'request', 'reserve', 'visit',
                'termin', 'cita', 'rdv', 'prenota',
            )):
                return True
            if any(p.search(href_l) for p in self.booking_url_patterns):
                return True
            # Vendor booking domains without login path still count (Zocdoc, Calendly…).
            return True

        return False

    def _has_booking_form(self, forms: List) -> bool:
        for form in forms:
            inputs = form.find_all('input')
            selects = form.find_all('select')
            textareas = form.find_all('textarea')
            labels = form.find_all('label')

            field_bits: List[str] = []
            has_password = False
            for field in inputs + selects + textareas:
                itype = (field.get('type') or '').lower()
                if itype == 'password':
                    has_password = True
                field_bits.extend([
                    (field.get('name') or '').lower(),
                    (field.get('placeholder') or '').lower(),
                    (field.get('id') or '').lower(),
                    ' '.join(field.get('class') or []).lower(),
                    (field.get('aria-label') or '').lower(),
                ])

            label_text = ' '.join(l.get_text(strip=True).lower() for l in labels)
            # Intent from structured fields/labels/action only — NOT full legal copy inside forms.
            form_action = (form.get('action') or '').lower()
            form_id = (form.get('id') or '').lower()
            form_class = ' '.join(form.get('class') or []).lower()
            form_name = (form.get('name') or '').lower()
            structured = ' '.join([
                ' '.join(field_bits), label_text, form_action, form_id, form_class, form_name,
            ])

            if has_password:
                continue
            if any(h in structured for h in NON_BOOKING_FORM_HINTS):
                # Still allow if action clearly appointment-related.
                if not any(tok in form_action for tok in (
                    'appointment', 'booking', 'schedule', 'termin', 'cita', 'rdv'
                )):
                    continue

            if any(tok in form_action for tok in (
                'appointment', 'booking', 'schedule', 'termin', 'cita', 'rdv', 'prenota'
            )):
                return True

            matched_categories = set()
            for category, patterns in APPOINTMENT_FORM_FIELD_PATTERNS.items():
                for pattern in patterns:
                    if re.search(pattern, structured, re.IGNORECASE):
                        matched_categories.add(category)
                        break

            # Strong appointment intent with word boundaries (avoid random "visit" in legalese).
            intent_re = re.compile(
                r'\b(appointment|appointments|booking|bookings|schedule|scheduling|'
                r'consultation|termin|rendez-?vous|\brdv\b|cita|prenota|afspraak)\b',
                re.I,
            )
            appointment_intent = bool(intent_re.search(structured))

            if appointment_intent and len(matched_categories) >= 2:
                return True

            if appointment_intent and matched_categories.intersection({'date', 'time', 'provider'}):
                return True

            date_inputs = form.find_all('input', {'type': 'date'})
            if date_inputs and appointment_intent:
                return True

            for select in selects:
                select_text = select.get_text(strip=True).lower()
                select_name = (
                    (select.get('name') or '') + ' ' +
                    (select.get('id') or '') + ' ' +
                    (select.get('aria-label') or '')
                ).lower()
                blob = f'{select_text} {select_name}'
                if appointment_intent and any(
                    word in blob for word in (
                        'date', 'time', 'doctor', 'provider', 'service', 'appointment',
                        'slot', 'termin',
                    )
                ):
                    return True

        return False

    def _has_contact_form_only(self, forms: List, text_content: str) -> bool:
        text_lower = text_content.lower()

        contact_indicators = sum(
            1 for kw in ('contact us', 'get in touch', 'reach us', 'send us a message')
            if kw in text_lower
        )
        appointment_indicators = sum(
            1 for kw in ('appointment', 'schedule', 'book online', 'termin', 'rendez-vous')
            if kw in text_lower
        )

        if contact_indicators > 0 and appointment_indicators == 0:
            return True

        for form in forms:
            inputs = form.find_all('input')
            selects = form.find_all('select')
            textareas = form.find_all('textarea')

            field_names = []
            for field in inputs + selects + textareas:
                field_names.append((field.get('name') or '').lower())
                field_names.append((field.get('placeholder') or '').lower())

            field_text = ' '.join(field_names)

            has_only_basic_contact = (
                bool(re.search(r'name', field_text)) and
                bool(re.search(r'email', field_text)) and
                bool(re.search(r'message|comment|text', field_text)) and
                not bool(re.search(
                    r'date|time|schedule|appointment|reason|provider|doctor|service|termin|cita',
                    field_text
                ))
            )
            if has_only_basic_contact:
                return True

        return False


def _uc_process_main(url: str, queue) -> None:
    """Child-process entry for SeleniumBase UC (must be top-level for pickle)."""
    try:
        from seleniumbase import SB

        html = ''
        with SB(uc=True, headless=True, test=False) as sb:
            try:
                sb.driver.set_page_load_timeout(20)
            except Exception:
                pass
            sb.uc_open_with_reconnect(url, reconnect_time=2)
            time.sleep(2.0)
            title = (sb.get_title() or '').lower()
            if (
                'just a moment' in title
                or 'attention required' in title
                or 'sgcaptcha' in (sb.get_current_url() or '').lower()
            ):
                time.sleep(4)
            html = sb.get_page_source() or ''
            # Reject hard block shells so callers don't score them as NO.
            if WebsiteFetcher.is_challenge_page(html):
                queue.put(('err', "Bot protection / captcha challenge"))
                return
            t = (sb.get_title() or '').lower()
            if 'forbidden' in t or t.strip().startswith('403'):
                queue.put(('err', "Bot protection / captcha challenge"))
                return
        queue.put(('ok', html))
    except Exception as e:
        queue.put(('err', f"UC render failed: {type(e).__name__}"))


class JsRenderer:
    """Playwright + SeleniumBase UC fallback for JS-heavy / Cloudflare pages.

    Order: Playwright (fast) → SeleniumBase UC (slower, clears many CF walls).
    UC runs at most once per site and is hard-capped so chromedriver cannot hang
    the whole scraper (seen on theintracare.com).
    """

    def __init__(self):
        self._available = None
        self._browser = None
        self._playwright = None
        self._uc_available = None
        self._uc_used_hosts: Set[str] = set()

    def available(self) -> bool:
        if self._available is not None:
            return self._available
        try:
            from playwright.sync_api import sync_playwright  # noqa: F401
            self._available = True
        except Exception:
            self._available = False
            logger.warning(
                "Playwright not available — JS-heavy pages may be under-detected. "
                "Install with: pip install playwright && playwright install chromium"
            )
        return self._available

    def uc_available(self) -> bool:
        if self._uc_available is not None:
            return self._uc_available
        try:
            from seleniumbase import SB  # noqa: F401
            self._uc_available = True
        except Exception:
            self._uc_available = False
        return self._uc_available

    def _ensure_browser(self):
        if self._browser is not None:
            return
        from playwright.sync_api import sync_playwright
        self._playwright = sync_playwright().start()
        self._browser = self._playwright.chromium.launch(
            headless=True,
            args=['--disable-blink-features=AutomationControlled'],
        )

    def _render_playwright(
        self, url: str, timeout_ms: int = JS_RENDER_TIMEOUT_MS
    ) -> Tuple[Optional[str], Optional[str]]:
        if not self.available():
            return None, "Playwright unavailable"
        try:
            self._ensure_browser()
            context = self._browser.new_context(
                user_agent=(
                    'Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                    'AppleWebKit/537.36 (KHTML, like Gecko) '
                    'Chrome/131.0.0.0 Safari/537.36'
                ),
                locale='en-US',
                viewport={'width': 1366, 'height': 768},
                ignore_https_errors=True,
            )
            page = context.new_page()
            page.add_init_script(
                "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
            )
            page.goto(url, wait_until='domcontentloaded', timeout=timeout_ms)
            time.sleep(min(JS_SETTLE_SECONDS, 2))
            html = page.content()
            current = (page.url or '').lower()
            context.close()

            if 'captcha' in current or 'sgcaptcha' in current or '__cf_chl' in current:
                return None, "Bot protection / captcha challenge"
            if WebsiteFetcher.is_challenge_page(html):
                return None, "Bot protection / captcha challenge"
            return html, None
        except Exception as e:
            return None, f"JS render failed: {type(e).__name__}"

    def _render_uc(self, url: str) -> Tuple[Optional[str], Optional[str]]:
        """SeleniumBase UC with a hard timeout (prevents infinite hangs)."""
        if not self.uc_available():
            return None, "UC browser unavailable"

        host = _strip_www(urlparse(url).netloc)
        if host in self._uc_used_hosts:
            return None, "Bot protection / captcha challenge"
        self._uc_used_hosts.add(host)

        logger.info(f"    Cloudflare UC bypass (max {UC_RENDER_TIMEOUT_SEC}s): {url}")

        import multiprocessing as mp

        ctx = mp.get_context('spawn')
        queue = ctx.Queue()
        proc = ctx.Process(target=_uc_process_main, args=(url, queue))
        proc.start()
        proc.join(timeout=UC_RENDER_TIMEOUT_SEC)

        if proc.is_alive():
            logger.warning(
                f"    UC bypass timed out after {UC_RENDER_TIMEOUT_SEC}s — killing browser"
            )
            proc.terminate()
            proc.join(timeout=5)
            if proc.is_alive():
                proc.kill()
                proc.join(timeout=2)
            return None, "Bot protection / captcha challenge"

        try:
            status, payload = queue.get_nowait()
        except Exception:
            return None, "Bot protection / captcha challenge"

        if status != 'ok':
            return None, payload
        html = payload
        if not html or WebsiteFetcher.is_challenge_page(html):
            return None, "Bot protection / captcha challenge"
        return html, None

    def render(
        self,
        url: str,
        timeout_ms: int = JS_RENDER_TIMEOUT_MS,
        allow_uc: bool = True,
    ) -> Tuple[Optional[str], Optional[str]]:
        """Return (html, error). Playwright first; UC only on captcha walls."""
        html, err = self._render_playwright(url, timeout_ms=timeout_ms)
        if html and not err:
            return html, None
        needs_uc = allow_uc and bool(err) and (
            'captcha' in err.lower()
            or 'bot protection' in err.lower()
            or 'timeout' in err.lower()
            or 'js render failed' in err.lower()
        )
        if needs_uc and self.uc_available():
            uc_html, uc_err = self._render_uc(url)
            if uc_html and not uc_err:
                return uc_html, None
            return None, uc_err or err
        return None, err

    def close(self):
        try:
            if self._browser:
                self._browser.close()
            if self._playwright:
                self._playwright.stop()
        except Exception:
            pass
        self._browser = None
        self._playwright = None
        self._uc_used_hosts.clear()


class LocalWorksheet:
    """Minimal gspread-like wrapper over a local .xlsx workbook."""

    def __init__(self, path: str, save_every: Optional[int] = None):
        self.path = path
        # keep_links=False reduces openpyxl memory on large sheets.
        self.wb = load_workbook(path, keep_links=False)
        self.ws = self.wb.active
        self._dirty = 0
        # Read LOCAL_SAVE_EVERY at runtime so --light can raise it.
        self._save_every = max(1, LOCAL_SAVE_EVERY if save_every is None else save_every)
        logger.info(
            f"Opened local workbook: {path} / tab: {self.ws.title!r} "
            f"({self.ws.max_row} rows)"
        )

    def row_values(self, row: int) -> List[str]:
        vals: List[str] = []
        for col in range(1, self.ws.max_column + 1):
            v = self.ws.cell(row=row, column=col).value
            vals.append('' if v is None else str(v))
        # Trim trailing empties to match gspread-ish behavior.
        while vals and not vals[-1].strip():
            vals.pop()
        return vals

    def col_values(self, col: int) -> List[str]:
        vals: List[str] = []
        for row in range(1, self.ws.max_row + 1):
            v = self.ws.cell(row=row, column=col).value
            vals.append('' if v is None else str(v))
        while vals and not vals[-1].strip():
            vals.pop()
        return vals

    def update_cell(self, row: int, col: int, value: str) -> None:
        self.ws.cell(row=row, column=col).value = value
        self._dirty += 1
        if self._dirty >= self._save_every:
            self.save()

    def save(self) -> None:
        self.wb.save(self.path)
        self._dirty = 0


def get_worksheet(path: str) -> LocalWorksheet:
    logger.info(f"Loading local workbook: {path}")
    return LocalWorksheet(path)


def resolve_columns(worksheet) -> Tuple[int, int]:
    """Find Website and Booking Feature columns by header name (1-based)."""
    headers = worksheet.row_values(1)
    headers_norm = [h.strip().lower() for h in headers]

    try:
        website_col = headers_norm.index(WEBSITE_HEADER.lower()) + 1
    except ValueError as exc:
        raise RuntimeError(
            f"Could not find '{WEBSITE_HEADER}' column in header row. Found: {headers}"
        ) from exc

    try:
        booking_col = headers_norm.index(BOOKING_HEADER.lower()) + 1
    except ValueError as exc:
        raise RuntimeError(
            f"Could not find '{BOOKING_HEADER}' column in header row. Found: {headers}"
        ) from exc

    logger.info(
        f"Columns resolved: {WEBSITE_HEADER}=col {website_col}, "
        f"{BOOKING_HEADER}=col {booking_col}"
    )
    return website_col, booking_col


def analyze_website(
    fetcher: WebsiteFetcher,
    detector: BookingDetector,
    website: str,
    js_renderer: Optional[JsRenderer] = None,
) -> WebsiteResult:
    """Inspect homepage then BFS-follow promising internal links up to MAX_PAGES_PER_SITE."""
    html, final_url, error = fetcher.fetch(website)
    used_js = False
    via_alternate = False

    def _js_url_variants(seed: str) -> List[str]:
        """Prefer www then bare host — many CDNs only allow the www form."""
        parsed = urlparse(seed)
        host = parsed.netloc
        if not host:
            return [seed]
        bare = _strip_www(host)
        path = parsed.path if parsed.path not in ('', '/') else ''
        variants = [
            f'https://www.{bare}{path}',
            f'https://{bare}{path}',
        ]
        out: List[str] = []
        seen: Set[str] = set()
        for v in variants:
            key = v.rstrip('/').lower()
            if key not in seen:
                seen.add(key)
                out.append(v.rstrip('/') or v)
        return out

    def _try_js(url: str) -> Tuple[Optional[str], Optional[str], Optional[str]]:
        if not (js_renderer and js_renderer.available()):
            return None, None, "Playwright unavailable"
        variants = _js_url_variants(url)
        last_js_err = None
        # 1) Fast Playwright on www + bare (no UC yet — avoids double Chrome hang).
        for candidate in variants:
            js_html, js_err = js_renderer.render(candidate, allow_uc=False)
            if js_html and not js_err and not WebsiteFetcher.is_challenge_page(js_html):
                return js_html, candidate, None
            last_js_err = js_err or "Bot protection / captcha challenge"
        # 2) One UC attempt on preferred URL only (hard-timeout inside render).
        preferred = variants[0] if variants else url
        js_html, js_err = js_renderer.render(preferred, allow_uc=True)
        if js_html and not js_err and not WebsiteFetcher.is_challenge_page(js_html):
            return js_html, preferred, None
        return None, None, js_err or last_js_err

    # Bot walls / captcha / connection failures: recover via JS and/or alternate URLs.
    err_low = (error or '').lower()
    hard_fetch_fail = bool(error) and (
        error.startswith('HTTP ')
        or 'captcha' in err_low
        or 'bot protection' in err_low
        or 'connection' in err_low
        or 'timed out' in err_low
        or 'timeout' in err_low
        or 'network is unreachable' in err_low
        or 'httperror' in err_low
        or 'ssl' in err_low
        or 'request failed' in err_low
        or 'max retries' in err_low
    )
    connectivity_fail = bool(error) and (
        'connection' in err_low
        or 'timed out' in err_low
        or 'timeout' in err_low
        or 'network is unreachable' in err_low
        or 'unreachable' in err_low
        or 'ssl' in err_low
        or 'request failed' in err_low
    )
    needs_browser = bool(error) and (
        'captcha' in err_low
        or 'bot protection' in err_low
        or error.startswith('HTTP 403')
        or error.startswith('HTTP 401')
        or error.startswith('HTTP 406')
        or error.startswith('HTTP 429')
        or error.startswith('HTTP 503')
        or 'httperror' in err_low
    )

    def _alt_is_usable(seed: str, alt: str, alt_html: str) -> bool:
        """Accept booking-platform URLs, or brand mirrors that actually book.

        Brand mirrors with no booking CTA are rejected so a captcha-blocked
        primary site does not get a false NO from an empty Wix/marketing clone.
        """
        if detector._matches_third_party(alt) or detector._is_booking_host(alt):
            return True
        hit, _ = detector.detect(alt_html, alt)
        return hit == Result.YES

    def _try_alternates(seed: str) -> bool:
        nonlocal html, final_url, error, via_alternate
        for alt in fetcher.discover_alternate_urls(seed, limit=3):
            alt_html, alt_final, alt_err = fetcher.fetch(alt)
            if alt_err or not alt_html or WebsiteFetcher.is_challenge_page(alt_html):
                continue
            if not _alt_is_usable(seed, alt, alt_html):
                continue
            html, final_url, error = alt_html, alt_final or alt, None
            via_alternate = True
            return True
        return False

    if hard_fetch_fail:
        norm = fetcher.normalize_url(website)
        if norm:
            recovered = False
            js_err = None
            # Browser first on 403/captcha — Chromium often loads pages that
            # curl/datacenter TLS cannot (e.g. PRC / Cloudflare sites).
            if needs_browser:
                js_html, js_url, js_err = _try_js(norm)
                if js_html:
                    html, final_url, error = js_html, js_url, None
                    used_js = True
                    recovered = True
            if not recovered:
                recovered = _try_alternates(norm)
            # Flaky connection hosts often load fine in a real browser after
            # curl gave up — try JS even when the failure looked like DNS/TCP.
            if not recovered and connectivity_fail and not used_js:
                js_html, js_url, js_err = _try_js(norm)
                if js_html:
                    html, final_url, error = js_html, js_url, None
                    used_js = True
                    recovered = True
            if not recovered and js_err and (needs_browser or connectivity_fail):
                error = js_err

    if error:
        return WebsiteResult(website, Result.ERROR, error)

    # Safety: never score a challenge shell as NO.
    if WebsiteFetcher.is_challenge_page(html):
        norm = fetcher.normalize_url(website) or final_url
        if norm and _try_alternates(norm):
            pass  # html/final_url recovered
        elif not used_js and norm:
            js_html, js_url, js_err = _try_js(norm)
            if js_html:
                html, final_url = js_html, js_url
                used_js = True
            else:
                return WebsiteResult(
                    website, Result.ERROR,
                    js_err or "Bot protection / captcha challenge",
                )
        else:
            return WebsiteResult(website, Result.ERROR, "Bot protection / captcha challenge")

    base_url = final_url or fetcher.normalize_url(website) or website
    visited: Set[str] = set()
    queue: deque = deque()

    def _norm(u: str) -> str:
        return (u or '').rstrip('/').lower()

    queue.append((base_url, html))
    visited.add(_norm(base_url))
    pages_checked = 0
    last_no_reason = "No booking feature detected"
    homepage_html = html

    while queue and pages_checked < MAX_PAGES_PER_SITE:
        if shutdown_handler.shutdown_requested:
            break

        page_url, page_html = queue.popleft()
        pages_checked += 1

        if WebsiteFetcher.is_challenge_page(page_html):
            continue

        page_result, page_reason = detector.detect(page_html, page_url)
        # Recovered onto a known scheduler URL (Zocdoc/Booker/ZipAppointments…)
        # even if the page is a thin SPA shell — that IS online booking.
        if (
            page_result != Result.YES
            and via_alternate
            and pages_checked == 1
            and (
                detector._matches_third_party(page_url)
                or detector._is_booking_host(page_url)
            )
        ):
            page_result = Result.YES
            page_reason = "Found third-party booking page"

        if page_result == Result.YES:
            bits = []
            if used_js:
                bits.append('JS')
            if via_alternate:
                bits.append(f'via {page_url}')
            elif pages_checked > 1:
                bits.append(f'via {page_url}')
            prefix = (bits[0] + ': ') if bits and bits[0] == 'JS' else ''
            suffix = ''
            if via_alternate or pages_checked > 1:
                suffix = f" (via {page_url})"
            return WebsiteResult(website, Result.YES, f"{prefix}{page_reason}{suffix}")

        last_no_reason = page_reason

        remaining = MAX_PAGES_PER_SITE - pages_checked - len(queue)
        if remaining <= 0:
            continue

        for follow_url in detector.collect_follow_urls(page_html, page_url, limit=remaining + 5):
            key = _norm(follow_url)
            if key in visited:
                continue
            visited.add(key)

            if PAGE_FETCH_DELAY:
                time.sleep(PAGE_FETCH_DELAY)

            follow_html, follow_final, follow_error = fetcher.fetch(follow_url)
            if follow_error or not follow_html or WebsiteFetcher.is_challenge_page(follow_html):
                continue

            follow_base = follow_final or follow_url
            visited.add(_norm(follow_base))
            queue.append((follow_base, follow_html))

            if pages_checked + len(queue) >= MAX_PAGES_PER_SITE:
                break

    # Selective Playwright fallback for JS-rendered booking UIs.
    if (
        js_renderer
        and js_renderer.available()
        and not used_js
        and (
            detector.looks_js_heavy(homepage_html)
            or detector.has_soft_booking_signal(homepage_html)
            or WebsiteFetcher.is_challenge_page(homepage_html)
        )
    ):
        js_html, js_err = js_renderer.render(base_url)
        if js_html and not js_err and not WebsiteFetcher.is_challenge_page(js_html):
            js_result, js_reason = detector.detect(js_html, base_url)
            if js_result == Result.YES:
                return WebsiteResult(website, Result.YES, f"JS: {js_reason}")

            for follow_url in detector.collect_follow_urls(
                js_html, base_url, limit=JS_FOLLOW_PAGES
            ):
                if shutdown_handler.shutdown_requested:
                    break
                fol_html, fol_err = js_renderer.render(follow_url)
                if fol_err or not fol_html or WebsiteFetcher.is_challenge_page(fol_html):
                    continue
                fol_result, fol_reason = detector.detect(fol_html, follow_url)
                if fol_result == Result.YES:
                    return WebsiteResult(
                        website, Result.YES, f"JS: {fol_reason} (via {follow_url})"
                    )

    return WebsiteResult(website, Result.NO, last_no_reason)


def process_websites(worksheet, website_col: int, booking_col: int):
    websites = worksheet.col_values(website_col)
    booking_values = worksheet.col_values(booking_col)

    total_rows = max(len(websites), 1)
    total_records = max(total_rows - 1, 0)

    fetcher = WebsiteFetcher()
    detector = BookingDetector()
    js_renderer: Optional[JsRenderer] = None
    if ENABLE_JS_RENDER:
        js_renderer = JsRenderer()
        if js_renderer.available():
            logger.info("Playwright JS fallback: enabled")
        else:
            logger.info("Playwright JS fallback: disabled")
        if js_renderer.uc_available():
            logger.info("Cloudflare UC fallback (SeleniumBase): enabled")
        else:
            logger.info(
                "Cloudflare UC fallback: disabled "
                "(pip install seleniumbase for better CF bypass)"
            )
    else:
        logger.info("Playwright/UC browsers: OFF (light mode)")

    processed = 0
    skipped_done = 0
    errors = 0

    try:
        for row_idx in range(2, total_rows + 1):
            if shutdown_handler.shutdown_requested:
                logger.info(f"Shutdown requested. Stopping at row {row_idx}")
                break

            website_raw = websites[row_idx - 1] if row_idx - 1 < len(websites) else ''
            website = (website_raw or '').strip()

            existing = ''
            if row_idx - 1 < len(booking_values):
                existing = (booking_values[row_idx - 1] or '').strip().upper()

            # Skip finished YES/NO/SKIP/ERROR (ERROR retry temporarily disabled).
            if existing in DONE_VALUES:
                skipped_done += 1
                continue

            if not website:
                logger.info(f"[{row_idx - 1}/{total_records}] Row {row_idx}: Empty website → SKIP")
                worksheet.update_cell(row_idx, booking_col, Result.SKIP.value)
                processed += 1
                continue

            normalized = fetcher.normalize_url(website)
            if not normalized:
                logger.info(f"[{row_idx - 1}/{total_records}] Checking: {website}")
                logger.info("    Booking Feature: SKIP (Invalid URL)")
                worksheet.update_cell(row_idx, booking_col, Result.SKIP.value)
                processed += 1
                continue

            logger.info(f"[{row_idx - 1}/{total_records}] Checking: {website}")

            outcome = analyze_website(fetcher, detector, website, js_renderer)
            logger.info(f"    Booking Feature: {outcome.result.value} ({outcome.reason})")
            worksheet.update_cell(row_idx, booking_col, outcome.result.value)

            if outcome.result == Result.ERROR:
                errors += 1
            processed += 1

            if processed % 100 == 0:
                logger.info(
                    f"Progress: {processed} written this run, "
                    f"{skipped_done} already done, {errors} errors"
                )

            time.sleep(REQUEST_DELAY)
    finally:
        if hasattr(worksheet, 'save'):
            worksheet.save()
            logger.info(f"Saved workbook: {worksheet.path}")
        if js_renderer is not None:
            js_renderer.close()

    logger.info(
        f"Completed: {processed} written, {skipped_done} already done, {errors} errors"
    )


def main():
    import argparse

    parser = argparse.ArgumentParser(
        description="Healthcare website booking-feature scraper (local .xlsx)."
    )
    parser.add_argument(
        "workbook",
        help="Path to the .xlsx file to scrape",
    )
    parser.add_argument(
        "--light",
        action="store_true",
        help=(
            "Low-RAM mode for parallel runs: no Playwright/UC browsers, "
            "fewer pages/retries. Prefer this when running multiple terminals."
        ),
    )
    args = parser.parse_args()
    workbook = args.workbook

    if args.light:
        apply_light_mode()

    logger.info("=" * 60)
    logger.info("Healthcare Website Booking Feature Scraper")
    logger.info("=" * 60)

    try:
        worksheet = get_worksheet(workbook)
        website_col, booking_col = resolve_columns(worksheet)

        booking_values = worksheet.col_values(booking_col)
        websites = worksheet.col_values(website_col)
        total = max(len(websites) - 1, 0)
        done = sum(
            1 for v in booking_values[1:]
            if (v or '').strip().upper() in DONE_VALUES
        )
        errors_pending = sum(
            1 for v in booking_values[1:]
            if (v or '').strip().upper() == 'ERROR'
        )
        remaining = total - done  # blanks only while ERROR is in DONE_VALUES

        logger.info(f"Local file: {workbook}")
        logger.info(
            f"Rows: {total} websites | already done: {done} | "
            f"remaining (blank): {remaining} "
            f"(ERROR parked, not retrying: {errors_pending})"
        )
        logger.info(
            f"Rate limiting: {REQUEST_DELAY}s between sites | "
            f"up to {MAX_PAGES_PER_SITE} pages/site | "
            f"JS render: {'on' if ENABLE_JS_RENDER else 'off'}"
        )
        logger.info("-" * 60)

        process_websites(worksheet, website_col, booking_col)

        logger.info("-" * 60)
        logger.info("Scraper finished — results are in the local .xlsx")
        logger.info(f"File: {workbook}")

    except FileNotFoundError:
        logger.error(
            f"Workbook not found: {workbook}. "
            "Check the path and retry."
        )
        raise
    except Exception as e:
        logger.error(f"Unexpected error: {type(e).__name__}: {e}")
        raise


if __name__ == "__main__":
    main()
